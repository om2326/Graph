import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette
    INK = RGBColor(24, 32, 44)         # #18202C
    BG_LIGHT = RGBColor(247, 246, 243) # #F7F6F3
    WHITE = RGBColor(255, 255, 255)
    TEAL = RGBColor(28, 110, 110)      # #1C6E6E
    RUST = RGBColor(181, 69, 45)       # #B5452D
    MUTED = RGBColor(100, 116, 139)    # #64748B
    ACCENT_BLUE = RGBColor(37, 99, 235)# #2563EB
    CARD_BG = RGBColor(255, 255, 255)
    LINE_COLOR = RGBColor(222, 218, 210)

    blank_layout = prs.slide_layouts[6]

    def add_header(slide, title_text, category_text="TECHNICAL ARCHITECTURE & DESIGN"):
        # Header banner
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(1.1))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p_cat = tf.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(10.5)
        p_cat.font.bold = True
        p_cat.font.color.rgb = TEAL

        p_title = tf.add_paragraph()
        p_title.text = title_text
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = INK

    def add_card(slide, left, top, width, height, title, items, badge=""):
        # Background card
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = CARD_BG
        shape.line.color.rgb = LINE_COLOR
        shape.line.width = Pt(1)

        tx_box = slide.shapes.add_textbox(left + Inches(0.25), top + Inches(0.2), width - Inches(0.5), height - Inches(0.4))
        tf = tx_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p_title = tf.paragraphs[0]
        p_title.text = title
        p_title.font.size = Pt(14)
        p_title.font.bold = True
        p_title.font.color.rgb = INK

        if badge:
            p_badge = tf.add_paragraph()
            p_badge.text = badge.upper()
            p_badge.font.size = Pt(8.5)
            p_badge.font.bold = True
            p_badge.font.color.rgb = TEAL
            p_badge.space_after = Pt(6)

        for item in items:
            p = tf.add_paragraph()
            p.text = f"•  {item}"
            p.font.size = Pt(11)
            p.font.color.rgb = INK
            p.space_before = Pt(4)
            p.space_after = Pt(2)

    # -------------------------------------------------------------
    # SLIDE 1: Title Slide
    # -------------------------------------------------------------
    slide1 = prs.slides.add_slide(blank_layout)
    bg = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = INK
    bg.line.fill.background()

    # Title box
    tbox = slide1.shapes.add_textbox(Inches(1.2), Inches(2.0), Inches(11.0), Inches(3.2))
    tf = tbox.text_frame
    tf.word_wrap = True

    p_sub0 = tf.paragraphs[0]
    p_sub0.text = "FULL-STACK ANALYTICS PLATFORM"
    p_sub0.font.size = Pt(13)
    p_sub0.font.bold = True
    p_sub0.font.color.rgb = RGBColor(45, 212, 191) # Light Teal
    p_sub0.space_after = Pt(12)

    p_main = tf.add_paragraph()
    p_main.text = "Correlation Network Studio"
    p_main.font.size = Pt(40)
    p_main.font.bold = True
    p_main.font.color.rgb = WHITE
    p_main.space_after = Pt(14)

    p_sub = tf.add_paragraph()
    p_sub.text = "End-to-End System Architecture, Statistical Engines, and Interactive Graph Visualizations"
    p_sub.font.size = Pt(18)
    p_sub.font.color.rgb = RGBColor(203, 213, 225)
    p_sub.space_after = Pt(24)

    p_meta = tf.add_paragraph()
    p_meta.text = "FastAPI  •  SQLite Database  •  D3.js Force Simulation  •  SciPy & Pandas"
    p_meta.font.size = Pt(12.5)
    p_meta.font.bold = True
    p_meta.font.color.rgb = RGBColor(148, 163, 184)

    # -------------------------------------------------------------
    # SLIDE 2: Problem Statement & Motivation
    # -------------------------------------------------------------
    slide2 = prs.slides.add_slide(blank_layout)
    add_header(slide2, "The Core Problem & Engineering Motivation", "Context & Objectives")

    add_card(slide2, Inches(0.8), Inches(1.8), Inches(3.7), Inches(4.8), 
             "The Tabular Bottleneck", 
             [
                 "Multivariate datasets contain rich hidden cross-feature dependencies.",
                 "Traditional correlation tables and matrices scale poorly as features grow.",
                 "Analysts struggle to visually detect clusters, community structures, or bridge variables from raw numbers.",
                 "Static spreadsheets fail to provide intuitive visual exploratory feedback."
             ], "Challenge")

    add_card(slide2, Inches(4.8), Inches(1.8), Inches(3.7), Inches(4.8), 
             "The Solution Approach", 
             [
                 "Transform numeric matrices into dynamic Force-Directed Graphs.",
                 "Nodes represent variables (e.g., student subjects, economic factors).",
                 "Edges represent statistically verified correlation links.",
                 "Edge weight, color, and line styling encode magnitude and sign directly.",
                 "Physics simulations organically organize clusters and central hubs."
             ], "Visual Analytics")

    add_card(slide2, Inches(8.8), Inches(1.8), Inches(3.7), Inches(4.8), 
             "Interactive Drilling", 
             [
                 "A high-level network graph alone is not enough—analysts need drill-down capability.",
                 "Clicking an edge reveals the underlying 2D scatter plot with linear regression.",
                 "Clicking a node inspects single-variable frequency distribution histograms.",
                 "Instant threshold adjustments filter out noise in real time."
             ], "Drill-Down")

    # -------------------------------------------------------------
    # SLIDE 3: Evolution: Prototype to Full-Stack
    # -------------------------------------------------------------
    slide3 = prs.slides.add_slide(blank_layout)
    add_header(slide3, "Project Evolution: From Prototype to Full-Stack", "Architecture Roadmap")

    add_card(slide3, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8), 
             "Phase 1: Client-Only Prototype", 
             [
                 "Initial file: correlation_network.html",
                 "Hardcoded synthetic dataset in client JavaScript (60 students).",
                 "In-browser calculation of Pearson correlation matrix.",
                 "Basic D3 force simulation with static threshold slider.",
                 "Limitations: No data persistence, no custom CSV upload, no statistical significance testing (p-values), no linked regression views."
             ], "Initial Prototype")

    add_card(slide3, Inches(6.9), Inches(1.8), Inches(5.6), Inches(4.8), 
             "Phase 2: Production Full-Stack Platform", 
             [
                 "Three-tier architecture with Backend, Frontend, and Database.",
                 "FastAPI backend with high-speed Pandas and SciPy calculations.",
                 "SQLite database for persistent datasets and record storage.",
                 "Advanced statistical rigor: Pearson & Spearman methods, p-value filtering.",
                 "Linked views: Interactive 2D scatter plots with OLS regression, distribution histograms.",
                 "CSV Ingestion: Upload and visualize any arbitrary numeric dataset.",
                 "Graph Ergonomics: Zoom/Pan (d3.zoom), node pinning, SVG export."
             ], "Current Architecture")

    # -------------------------------------------------------------
    # SLIDE 4: Three-Tier Architecture
    # -------------------------------------------------------------
    slide4 = prs.slides.add_slide(blank_layout)
    add_header(slide4, "Full-Stack System Architecture", "System Design")

    add_card(slide4, Inches(0.8), Inches(1.8), Inches(3.7), Inches(4.8), 
             "1. Frontend Tier", 
             [
                 "Modular JavaScript architecture (api.js, graph.js, heatmap.js, charts.js, app.js).",
                 "D3.js v7 interactive physics force simulation.",
                 "Dual views: Network Graph & Correlation Heatmap Matrix.",
                 "Zoom, pan, node dragging, and double-click pinning.",
                 "Modals for OLS scatter plots and custom CSV uploads.",
                 "Accessible CSS styling with positive/negative color scales."
             ], "Client Side")

    add_card(slide4, Inches(4.8), Inches(1.8), Inches(3.7), Inches(4.8), 
             "2. Backend Tier", 
             [
                 "FastAPI application (backend/app.py) running on Uvicorn.",
                 "Stats engine (backend/stats.py) powered by SciPy and NumPy.",
                 "Calculates correlation coefficients, two-tailed p-values, linear regression (slope, intercept, R²), and histogram binning.",
                 "Multipart CSV upload parsing and schema validation.",
                 "Static asset hosting and auto-generated Swagger documentation."
             ], "Server Side")

    add_card(slide4, Inches(8.8), Inches(1.8), Inches(3.7), Inches(4.8), 
             "3. Database Tier", 
             [
                 "SQLite relational database (data/app.db).",
                 "datasets table: Stores metadata, column lists, and row counts.",
                 "records table: High-performance storage of dataset rows.",
                 "Database seeder (backend/seed_data.py) preloading realistic benchmarks.",
                 "Zero-configuration deployment, lightweight, and ACID-compliant."
             ], "Persistence")

    # -------------------------------------------------------------
    # SLIDE 5: Mathematical & Statistical Core
    # -------------------------------------------------------------
    slide5 = prs.slides.add_slide(blank_layout)
    add_header(slide5, "Mathematical & Statistical Engines", "Algorithms & Mathematics")

    add_card(slide5, Inches(0.8), Inches(1.8), Inches(5.6), Inches(2.3), 
             "Correlation Coefficient Formulations", 
             [
                 "Pearson Correlation (r): Measures linear co-variation between variables X and Y normalized by standard deviations: r = cov(X,Y) / (σx · σy).",
                 "Spearman Rank Correlation (ρ): Non-parametric measure of monotonic relationships based on rank orders: ρ = 1 - (6 Σ d²) / (n(n² - 1))."
             ], "Correlation Metrics")

    add_card(slide5, Inches(6.9), Inches(1.8), Inches(5.6), Inches(2.3), 
             "Statistical Significance & p-Values", 
             [
                 "Prevents spurious correlation artifacts from small sample sizes.",
                 "Computes two-tailed Student t-distribution test: t = r · sqrt((n-2) / (1 - r²)).",
                 "Users can toggle a strict filter (p < 0.05) to only display statistically significant dependencies."
             ], "Hypothesis Testing")

    add_card(slide5, Inches(0.8), Inches(4.3), Inches(5.6), Inches(2.3), 
             "Ordinary Least Squares (OLS) Regression", 
             [
                 "For any selected pair of variables, fits linear model y = β₁x + β₀.",
                 "Calculates coefficient of determination R² representing proportion of explained variance.",
                 "Evaluates slope, intercept, and standard error."
             ], "Regression Engine")

    add_card(slide5, Inches(6.9), Inches(4.3), Inches(5.6), Inches(2.3), 
             "Frequency Distribution Binning", 
             [
                 "Calculates NumPy frequency histograms across dynamic bin ranges.",
                 "Computes distribution metrics: Mean, Standard Deviation, Median, Min, Max, and Interquartile Ranges (IQR)."
             ], "Distributions")

    # -------------------------------------------------------------
    # SLIDE 6: Force-Directed Graph Mechanics
    # -------------------------------------------------------------
    slide6 = prs.slides.add_slide(blank_layout)
    add_header(slide6, "D3.js Force Simulation Mechanics", "Data Visualization")

    add_card(slide6, Inches(0.8), Inches(1.8), Inches(3.7), Inches(4.8), 
             "Physics Simulation Forces", 
             [
                 "Charge Force (ManyBody): Electrostatic repulsion (-320 strength) pushing nodes apart to prevent visual clutter.",
                 "Center Force: Pulls active graph centroid toward the SVG origin.",
                 "Collision Force: Radius boundary (45px) preventing overlapping nodes and labels.",
                 "Link Force: Spring tension that attracts correlated nodes with distance = 200 - |r|·120."
             ], "Force Engine")

    add_card(slide6, Inches(4.8), Inches(1.8), Inches(3.7), Inches(4.8), 
             "Visual Encoding Rules", 
             [
                 "Node Size: Scales with Degree Strength (sum of all connected |r| values). Central variables appear prominent.",
                 "Edge Color: Teal (#1C6E6E) indicates positive correlation; Rust (#B5452D) indicates inverse correlation.",
                 "Edge Texture: Solid lines for positive, dashed lines for negative (accessible for colorblindness).",
                 "Edge Width: Proportional to correlation magnitude."
             ], "Encoding System")

    add_card(slide6, Inches(8.8), Inches(1.8), Inches(3.7), Inches(4.8), 
             "User Ergonomics", 
             [
                 "Smooth Zoom & Pan (d3.zoom) with reset controls.",
                 "Interactive Drag: Drag any node to inspect local topology.",
                 "Double-Click Pinning: Lock node coordinates (fx, fy) to preserve custom analytical layouts.",
                 "Contextual Dimming: Clicking a node isolates its immediate neighborhood and dims unrelated elements."
             ], "Interaction")

    # -------------------------------------------------------------
    # SLIDE 7: Linked Views: Scatter Plots & Heatmap
    # -------------------------------------------------------------
    slide7 = prs.slides.add_slide(blank_layout)
    add_header(slide7, "Multi-View Linked Analytics", "Analytical Workflow")

    add_card(slide7, Inches(0.8), Inches(1.8), Inches(3.7), Inches(4.8), 
             "Correlation Heatmap Matrix", 
             [
                 "Switch instantly between Network Graph and N × N symmetric heatmap grid.",
                 "Diverging color gradient: Dark Rust (-1.0) -> Neutral White (0.0) -> Dark Teal (+1.0).",
                 "Displays numerical coefficients inside cells when zoomed.",
                 "Click any matrix cell to open its corresponding scatter plot."
             ], "Heatmap View")

    add_card(slide7, Inches(4.8), Inches(1.8), Inches(3.7), Inches(4.8), 
             "Scatter Plot Modal", 
             [
                 "Clicking any edge in the graph or cell in the heatmap launches the 2D scatter view.",
                 "Plots individual observations (e.g. students or periods).",
                 "Draws fitted linear regression trendline.",
                 "Displays equation, Pearson r, R², and p-value readout."
             ], "Bivariate Inspection")

    add_card(slide7, Inches(8.8), Inches(1.8), Inches(3.7), Inches(4.8), 
             "Inspector & Histograms", 
             [
                 "Clicking any node opens the side Inspector card.",
                 "Generates frequency distribution histogram of the selected variable.",
                 "Renders key statistical summary: Mean, Std Dev, Min, Max, Active Link Count.",
                 "Lists connected neighbors ranked by correlation strength."
             ], "Univariate Inspection")

    # -------------------------------------------------------------
    # SLIDE 8: Data Ingestion & Seeding
    # -------------------------------------------------------------
    slide8 = prs.slides.add_slide(blank_layout)
    add_header(slide8, "Data Ingestion & Benchmark Datasets", "Data Management")

    add_card(slide8, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8), 
             "Pre-Seeded Benchmark Datasets", 
             [
                 "Dataset 1: Student Academic Performance (N=100)",
                 "• 8 subjects: Math, Physics, Chemistry, CS, English, History, Economics, Biology.",
                 "• Generated using latent factor modeling (Quantitative, Verbal, Memory) with Box-Muller Gaussian noise.",
                 "• Realistically replicates STEM vs. Humanities disciplinary clusters.",
                 "",
                 "Dataset 2: Macroeconomic Indicators (N=80)",
                 "• 6 variables: GDP Growth, Inflation, Unemployment, Interest Rates, Consumer Confidence, Stock Index.",
                 "• Demonstrates strong inverse correlations (e.g., GDP vs. Unemployment r = -0.95)."
             ], "Preloaded Data")

    add_card(slide8, Inches(6.9), Inches(1.8), Inches(5.6), Inches(4.8), 
             "Custom CSV Ingestion Pipeline", 
             [
                 "Direct file upload via modal dialog.",
                 "Backend validation using Pandas: automatically isolates numeric columns and handles missing values.",
                 "Requires at least 2 numeric features and 3 data rows.",
                 "Generates unique dataset UUID and stores rows in SQLite.",
                 "Immediately populates dropdown and renders graph without server restart."
             ], "File Upload Engine")

    # -------------------------------------------------------------
    # SLIDE 9: REST API Specification
    # -------------------------------------------------------------
    slide9 = prs.slides.add_slide(blank_layout)
    add_header(slide9, "Backend REST API Architecture", "API Endpoints")

    add_card(slide9, Inches(0.8), Inches(1.8), Inches(5.6), Inches(2.3), 
             "GET /api/datasets", 
             [
                 "Returns list of all available datasets (metadata, row counts, column names, default flags).",
                 "Response: JSON array of dataset objects."
             ], "Dataset Discovery")

    add_card(slide9, Inches(6.9), Inches(1.8), Inches(5.6), Inches(2.3), 
             "GET /api/datasets/{id}/correlations", 
             [
                 "Query Params: method (pearson/spearman), min_threshold (0.0-1.0), max_p_value (1.0 or 0.05).",
                 "Returns active nodes, links, full matrix, and summary statistics."
             ], "Graph Calculation")

    add_card(slide9, Inches(0.8), Inches(4.3), Inches(5.6), Inches(2.3), 
             "GET /api/datasets/{id}/scatter", 
             [
                 "Query Params: x, y (column names).",
                 "Returns: Array of {x, y} coordinate points and linear regression results (slope, intercept, R², p-value)."
             ], "Bivariate Regression")

    add_card(slide9, Inches(6.9), Inches(4.3), Inches(5.6), Inches(2.3), 
             "POST /api/datasets/upload", 
             [
                 "Form Data: file (.csv), name, description.",
                 "Parses CSV, validates numeric features, inserts rows into SQLite, and returns new dataset ID."
             ], "Data Upload")

    # -------------------------------------------------------------
    # SLIDE 10: Technical Stack Summary
    # -------------------------------------------------------------
    slide10 = prs.slides.add_slide(blank_layout)
    add_header(slide10, "Technical Stack & Implementation Summary", "Technology Stack")

    add_card(slide10, Inches(0.8), Inches(1.8), Inches(2.8), Inches(4.8), 
             "Frontend", 
             [
                 "HTML5 & CSS3",
                 "D3.js v7 (Force, Zoom, Scales, Selection)",
                 "Modular ES6 JavaScript",
                 "SVG Vector Graphics",
                 "Responsive CSS Grid & Flexbox"
             ], "Client Tier")

    add_card(slide10, Inches(3.8), Inches(1.8), Inches(2.8), Inches(4.8), 
             "Backend", 
             [
                 "Python 3.14",
                 "FastAPI 0.138",
                 "Uvicorn ASGI Server",
                 "SciPy 1.18 (Stats & Regression)",
                 "Pandas 3.0 & NumPy 2.4"
             ], "Server Tier")

    add_card(slide10, Inches(6.8), Inches(1.8), Inches(2.8), Inches(4.8), 
             "Database", 
             [
                 "SQLite Relational Engine",
                 "ACID Compliant",
                 "JSON Column Serialization",
                 "Indexed foreign keys",
                 "Zero-configuration setup"
             ], "Database Tier")

    add_card(slide10, Inches(9.8), Inches(1.8), Inches(2.8), Inches(4.8), 
             "Dev & Deployment", 
             [
                 "Single-command launch (python run.py)",
                 "Windows Batch launcher (start.bat)",
                 "Auto-reloading dev server",
                 "Swagger Docs (/docs)"
             ], "Tooling")

    # -------------------------------------------------------------
    # SLIDE 11: Summary & Demonstration
    # -------------------------------------------------------------
    slide11 = prs.slides.add_slide(blank_layout)
    add_header(slide11, "Project Summary & Live Demonstration", "Conclusion")

    add_card(slide11, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8), 
             "Key Achievements", 
             [
                 "Successfully transitioned from a static single-file script to an enterprise-grade full-stack application.",
                 "High-performance statistical computing in Python integrated seamlessly with interactive D3.js physics.",
                 "Combines high-level network topology with granular bivariate scatter regression and univariate histograms.",
                 "Supports arbitrary user CSV datasets with zero manual schema configuration.",
                 "Engineered with intuitive Windows one-click launching."
             ], "Outcomes")

    add_card(slide11, Inches(6.9), Inches(1.8), Inches(5.6), Inches(4.8), 
             "How to Run the Live System", 
             [
                 "1. Double-click start.bat in the project directory, OR",
                 "2. Run from terminal: python run.py",
                 "3. Web Dashboard: http://127.0.0.1:8000",
                 "4. API Swagger Documentation: http://127.0.0.1:8000/docs",
                 "",
                 "Thank you! Ready for questions & demonstration."
             ], "Demonstration")

    out_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "presentation.pptx")
    prs.save(out_path)
    print(f"Presentation saved successfully to: {out_path}")

if __name__ == "__main__":
    create_deck()
